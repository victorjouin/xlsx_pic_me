"""
xlsxtiles.render — du plan de tuiles aux fichiers PNG.

Chaîne : openpyxl (met en scène une feuille : print_area, titres répétés,
sauts de page, taille de papier) -> soffice --convert-to pdf -> PyMuPDF
(rognage au contenu, rasterisation, vérification de grille).

Deux entrées :
  - render_sheet_tiles() : tuilage complet d'une feuille ;
  - render_citation()    : fenêtre à la demande autour de lignes citées.
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import openpyxl
import pymupdf

from .config import (
    DEFAULT_PATCH_BUDGET, GRID_DRIFT_TOLERANCE_PX, PAGE_SAFETY_H, PAGE_SAFETY_W,
    RENDER_DPI, TRIM_PAD_PT,
)
from .grid import SheetGrid, build_grid, content_clip, verify_grid
from .models import RenderedTile, TilePlan, WorkbookProfile
from .plan import apply_display_widths, compute_display_widths, plan_tiles
from .refs import col_letters
from .soffice import recalc_with_libreoffice, seed_profile, to_pdf

__all__ = ["render_sheet_tiles", "render_citation", "safe_name"]


def safe_name(name: str) -> str:
    """Nom de feuille -> fragment de nom de fichier utilisable sur tout OS."""
    out = "".join(c if c.isalnum() or c in "-_" else "_" for c in name)
    return (out.strip("_") or "sheet")[:40]


# --------------------------------------------------------------------------
# Mise en scène du classeur
# --------------------------------------------------------------------------

def prepare_workbook(src: Path, sheet_name: str, plans: list[TilePlan],
                     dest: Path, grid: SheetGrid | None = None,
                     display_widths: dict[int, float] | None = None,
                     extra_width_px: float = 0.0) -> None:
    """Écrit un classeur ne contenant que la feuille voulue, paginée sur le plan.

    Une seule conversion soffice pour les N tuiles d'une même bande de colonnes :
    on pose les sauts de ligne correspondant au plan et on fixe l'ordre de
    pagination, ce qui rend le mapping page -> tuile déterministe. Appeler
    soffice par tuile serait ~2 s x N.

    `extra_width_px` : largeur des colonnes clés répétées par LibreOffice sur
    une bande qui ne commence pas à la première colonne de la feuille. Sans
    elle la page est trop étroite et la bande déborde sur une page de plus.
    """
    wb = openpyxl.load_workbook(src)
    for name in list(wb.sheetnames):
        if name != sheet_name:
            del wb[name]
    ws = wb[sheet_name]

    r0 = min(p.row_start for p in plans)
    r1 = max(p.row_end for p in plans)
    c0 = min(p.col_start for p in plans)
    c1 = max(p.col_end for p in plans)

    for col, w in (display_widths or {}).items():
        ws.column_dimensions[col_letters(col)].width = w

    ws.print_area = f"{col_letters(c0)}{r0}:{col_letters(c1)}{r1}"
    if plans[0].header_rows:
        ws.print_title_rows = plans[0].header_rows
    if plans[0].key_cols:
        ws.print_title_cols = plans[0].key_cols

    ws.page_setup.scale = 100          # PAS fitToPage : il casse le mapping pixel
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.pageOrder = "overThenDown"
    ws.print_options.gridLines = True
    ws.print_options.headings = False

    margin_in = 0.08
    ws.page_margins.left = ws.page_margins.right = margin_in
    ws.page_margins.top = ws.page_margins.bottom = margin_in
    ws.page_margins.header = ws.page_margins.footer = 0.0

    # Page dimensionnée SUR LA TUILE. Sans ça, LibreOffice garde le format par
    # défaut (Letter/A4), repagine selon sa propre logique, et le mapping
    # page -> tuile saute : test mesuré à 2 pages PDF pour 1 tuile planifiée.
    if grid is not None:
        scale = RENDER_DPI / 96.0
        w_in = (grid.width_px(c0, c1) + extra_width_px) / scale / 96.0
        h_in = max(grid.height_px(p.row_start, p.row_end) for p in plans) / scale / 96.0
        if plans[0].header_rows:
            hr0, hr1 = (int(x) for x in plans[0].header_rows.split(":"))
            h_in += grid.height_px(hr0, hr1) / scale / 96.0
        # Marge de sécurité GÉNÉREUSE, et c'est délibéré. La grille prédite
        # n'égale jamais exactement le rendu LibreOffice (arrondis de largeur
        # par colonne, cumulés sur toute la bande) : mesuré ici, 7 colonnes
        # suffisent à faire déborder une tuile sur une 2e page avec 2 % de
        # marge. Or une page trop grande ne coûte RIEN — content_clip() rogne
        # le blanc avant rasterisation. Une page trop petite, elle, casse le
        # mapping page -> tuile. On surdimensionne donc franchement.
        # Les sauts de page manuels posés plus bas empêchent deux tuiles de
        # se retrouver sur la même page malgré la hauteur excédentaire.
        w_mm = (w_in + 2 * margin_in) * 25.4 * PAGE_SAFETY_W
        h_mm = (h_in + 2 * margin_in) * 25.4 * PAGE_SAFETY_H
        ws.page_setup.paperWidth = f"{w_mm:.2f}mm"
        ws.page_setup.paperHeight = f"{h_mm:.2f}mm"
        ws.page_setup.paperSize = None

    # openpyxl : row_breaks est un objet RowBreak, pas une liste ; la liste
    # des sauts est dans .brk
    from openpyxl.worksheet.pagebreak import Break
    ws.row_breaks.brk = []
    ws.col_breaks.brk = []
    for b in sorted({p.row_end for p in plans if p.row_end < r1}):
        ws.row_breaks.append(Break(id=b))
    for b in sorted({p.col_end for p in plans if p.col_end < c1}):
        ws.col_breaks.append(Break(id=b))

    wb.save(dest)


# --------------------------------------------------------------------------
# Rendu d'une bande de colonnes
# --------------------------------------------------------------------------

def _render_band(work: Path, sheet_name: str, band: list[TilePlan], grid: SheetGrid,
                 widths: dict[int, float], extra_w: float, outdir: Path,
                 tmp: Path, prof_dir: Path, calibration_ok: bool,
                 dpi: int, tag: str) -> list[RenderedTile]:
    """Convertit une bande de colonnes en une passe et découpe ses pages.

    Retourne une liste VIDE si la pagination ne correspond pas au plan :
    l'appelant reprend alors tuile par tuile. On ne devine jamais quelle page
    porte quelle plage — livrer une image dont on ignore la plage est pire que
    payer N conversions.
    """
    staged = tmp / f"stage_{tag}.xlsx"
    prepare_workbook(work, sheet_name, band, staged, grid, widths, extra_w)
    pdf_path = to_pdf(staged, tmp, prof_dir)

    tiles: list[RenderedTile] = []
    with pymupdf.open(pdf_path) as doc:
        if len(doc) != len(band):
            return []
        for i, page in enumerate(doc):
            pix = page.get_pixmap(dpi=dpi, clip=content_clip(page, TRIM_PAD_PT))
            png = outdir / f"{safe_name(sheet_name)}_{band[i].index:03d}.png"
            pix.save(png)
            drift, n_anchor = verify_grid(
                page, grid, band[i].col_start, band[i].col_end, dpi)
            tiles.append(RenderedTile(
                plan=band[i], png_path=str(png),
                width_px=pix.width, height_px=pix.height,
                grid_verified=calibration_ok and drift <= GRID_DRIFT_TOLERANCE_PX,
                grid_drift_px=round(drift, 2), n_anchor_checks=n_anchor,
            ))
    return tiles


def render_sheet_tiles(src: str | Path, profile: WorkbookProfile, sheet_name: str,
                       outdir: str | Path, *,
                       patch_budget: int = DEFAULT_PATCH_BUDGET,
                       dpi: int = RENDER_DPI) -> list[RenderedTile]:
    """Rend une feuille en tuiles PNG cadrées sur leur contenu.

    Une conversion soffice par bande de colonnes. Une bande dont la pagination
    ne correspond pas au plan est reprise tuile par tuile.
    """
    src, outdir = Path(src), Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    sp = profile.sheet(sheet_name)
    if sp is None:
        raise ValueError(f"feuille inconnue : {sheet_name}")

    merged_widths = apply_display_widths(
        sp, compute_display_widths(src, sheet_name, sp))
    grid = build_grid(sp, profile.calibration)
    plans = plan_tiles(sp, grid, patch_budget=patch_budget)
    cal_ok = bool(profile.calibration.calibrated)

    # bandes de colonnes, dans l'ordre du plan
    bands: dict[tuple[int, int], list[TilePlan]] = {}
    for p in plans:
        bands.setdefault((p.col_start, p.col_end), []).append(p)

    c0 = sp.true_min_col or 1
    tiles: list[RenderedTile] = []
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        prof_dir = seed_profile(tmp / "loprofile")
        work = src
        if profile.verdict.needs_recalc:
            work = recalc_with_libreoffice(src, tmp, prof_dir)

        for bi, ((bc0, _bc1), band) in enumerate(bands.items()):
            # colonnes clés répétées : elles s'ajoutent à la largeur de page
            extra_w = 0.0
            if band[0].key_cols and bc0 > c0 and sp.freeze_cols:
                extra_w = grid.width_px(c0, c0 + sp.freeze_cols - 1)

            got = _render_band(work, sheet_name, band, grid, merged_widths,
                               extra_w, outdir, tmp, prof_dir, cal_ok, dpi,
                               f"{safe_name(sheet_name)}_b{bi}")
            if not got:
                for p in band:
                    got.extend(_render_band(
                        work, sheet_name, [p], grid, merged_widths, extra_w,
                        outdir, tmp, prof_dir, cal_ok, dpi,
                        f"{safe_name(sheet_name)}_b{bi}_t{p.index}"))
            tiles.extend(got)

    tiles.sort(key=lambda t: t.plan.index)
    return tiles


def render_citation(src: str | Path, profile: WorkbookProfile, sheet_name: str,
                    row_start: int, row_end: int, outdir: str | Path, *,
                    context_rows: int = 3, dpi: int = RENDER_DPI) -> RenderedTile:
    """Rend une fenêtre autour des lignes citées par l'agent.

    Complémentaire du tuilage : sert à produire, à la demande, l'image exacte
    des lignes qu'une réponse cite, avec quelques lignes de contexte.
    """
    sp = profile.sheet(sheet_name)
    if sp is None:
        raise ValueError(f"feuille inconnue : {sheet_name}")

    header_n = max((t.header_row_count for t in sp.tables), default=sp.freeze_rows)
    r0 = max((sp.true_min_row or 1) + header_n, row_start - context_rows)
    r1 = min(sp.true_max_row or row_end, row_end + context_rows)
    c0, c1 = sp.true_min_col or 1, sp.true_max_col or 1

    merged_widths = apply_display_widths(
        sp, compute_display_widths(Path(src), sheet_name, sp))
    grid = build_grid(sp, profile.calibration)
    plan = TilePlan(0, r0, r1, c0, c1,
                    f"{sp.true_min_row}:{(sp.true_min_row or 1) + header_n - 1}"
                    if header_n else None,
                    None, f"{col_letters(c0)}{r0}:{col_letters(c1)}{r1}")

    src, outdir = Path(src), Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        prof_dir = seed_profile(tmp / "loprofile")
        staged = tmp / "cite.xlsx"
        prepare_workbook(src, sheet_name, [plan], staged, grid, merged_widths)
        pdf_path = to_pdf(staged, tmp, prof_dir)
        with pymupdf.open(pdf_path) as doc:
            page = doc[0]
            pix = page.get_pixmap(dpi=dpi, clip=content_clip(page, TRIM_PAD_PT))
            png = outdir / f"cite_{safe_name(sheet_name)}_{r0}-{r1}.png"
            pix.save(png)
            drift, anchors = verify_grid(page, grid, c0, c1, dpi)
            return RenderedTile(plan, str(png), pix.width, pix.height,
                                bool(profile.calibration.calibrated)
                                and drift <= GRID_DRIFT_TOLERANCE_PX,
                                round(drift, 2), anchors)
